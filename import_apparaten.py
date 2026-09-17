"""
Parser en validatie voor het importeren van een apparatenbestand (.xlsx of .ods).

Bedoeld voor twee gebruiksscenario's:
  1. Round-trip: exporteer met "Exporteer naar Excel/ODS", bewerk in Excel/
     LibreOffice/Numbers, importeer weer terug.
  2. Een handmatig voorbereid bestand met (een subset van) dezelfde
     kolomnamen als het apparatenregister.

Kolomnamen worden hoofdletterongevoelig gematcht. Alleen 'nummer' en 'klasse'
zijn verplicht; de rest is optioneel en krijgt een lege waarde als hij ontbreekt.
Bevat geen enkele schrijf-logica naar apparaten.csv - dat doet app.py, na
bevestiging door de gebruiker (zie /api/import/apparaten/analyseer en
/api/import/apparaten/bevestig).
"""

from pathlib import Path
from typing import Optional

GELDIGE_KLASSEN = {"I", "II", "snoer"}
NUMERIEKE_VELDEN = ("keuringstermijn_maanden", "kabellengte_m", "kabeldiameter_mm2")


def _cel_naar_tekst(waarde) -> str:
    if waarde is None:
        return ""
    return str(waarde).strip()


def _lees_xlsx(pad: Path) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(pad, data_only=True)
    ws = wb["apparaten-export"] if "apparaten-export" in wb.sheetnames else wb[wb.sheetnames[0]]

    rijen_ruw = list(ws.iter_rows(values_only=True))
    if not rijen_ruw:
        return []
    header = [_cel_naar_tekst(h).lower() for h in rijen_ruw[0]]

    resultaat = []
    for rij in rijen_ruw[1:]:
        if all(c is None or _cel_naar_tekst(c) == "" for c in rij):
            continue
        resultaat.append({header[i]: _cel_naar_tekst(rij[i]) for i in range(len(header)) if i < len(rij)})
    return resultaat


def _lees_ods(pad: Path) -> list[dict]:
    from odf.opendocument import load
    from odf.table import Table, TableRow, TableCell
    from odf import teletype

    doc = load(str(pad))
    tabellen = doc.spreadsheet.getElementsByType(Table)
    tabel = next((t for t in tabellen if t.getAttribute("name") == "apparaten-export"), tabellen[0])

    alle_rijen = []
    for rij in tabel.getElementsByType(TableRow):
        waarden = []
        for cell in rij.getElementsByType(TableCell):
            tekst = teletype.extractText(cell)
            herhaal = int(cell.getAttribute("numbercolumnsrepeated") or 1)
            waarden.extend([tekst] * herhaal)
        alle_rijen.append(waarden)

    if not alle_rijen:
        return []
    header = [_cel_naar_tekst(h).lower() for h in alle_rijen[0]]

    resultaat = []
    for rij in alle_rijen[1:]:
        if all(_cel_naar_tekst(c) == "" for c in rij):
            continue
        resultaat.append({header[i]: _cel_naar_tekst(rij[i]) for i in range(len(header)) if i < len(rij)})
    return resultaat


def parse_apparatenbestand(pad) -> list[dict]:
    pad = Path(pad)
    if pad.suffix.lower() == ".ods":
        return _lees_ods(pad)
    return _lees_xlsx(pad)


def valideer_rij(rij: dict, index: int, apparaat_velden: list[str]) -> dict:
    """Normaliseert een ruwe rij en controleert hem. Geeft altijd terug:
    {index, fouten: [...], genormaliseerd: {...}} - genormaliseerd bevat
    elk veld uit apparaat_velden (leeg als het ontbrak)."""
    fouten = []
    genormaliseerd = {veld: rij.get(veld, "") for veld in apparaat_velden}

    nummer = genormaliseerd.get("nummer", "").strip()
    genormaliseerd["nummer"] = nummer
    if not nummer:
        fouten.append("nummer ontbreekt")

    klasse = genormaliseerd.get("klasse", "").strip()
    genormaliseerd["klasse"] = klasse
    if klasse not in GELDIGE_KLASSEN:
        fouten.append(f"klasse '{klasse}' is ongeldig (verwacht: I, II of snoer)")

    for numveld in NUMERIEKE_VELDEN:
        waarde = genormaliseerd.get(numveld, "")
        if waarde:
            try:
                float(str(waarde).replace(",", "."))
            except ValueError:
                fouten.append(f"{numveld} ('{waarde}') is geen geldig getal")

    return {"index": index, "fouten": fouten, "genormaliseerd": genormaliseerd}